from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.management import call_command
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.audit.models import AuditEntry
from apps.donors.models import Donor
from .forms import DonationForm
from .models import Donation


class ImprovementsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('setup_roles', stdout=StringIO())
        cls.admin = get_user_model().objects.create_user(username='admin')
        cls.operator = get_user_model().objects.create_user(username='operator')
        cls.reader = get_user_model().objects.create_user(username='reader')
        cls.unassigned = get_user_model().objects.create_user(username='unassigned')
        for user, group in [(cls.admin, 'Administrador'), (cls.operator, 'Operador'), (cls.reader, 'Consulta')]:
            user.groups.add(Group.objects.get(name=group))
        cls.donor = Donor.objects.create(name='Maria & José', cpf='52998224725')
        cls.donation = Donation.objects.create(
            donor=cls.donor, donation_date=timezone.localdate(),
            amount=Decimal('1234.56'), payment_method='pix', created_by=cls.operator,
        )

    def setUp(self):
        self.client.force_login(self.admin)

    def url(self, name):
        return reverse(f'donations:{name}', args=[self.donation.pk])

    def donation_payload(self, **overrides):
        payload = {'donor': self.donor.pk, 'donation_date': timezone.localdate().isoformat(), 'donation_amount': '50.00', 'contribution_amount': '0.00', 'payment_method': 'pix'}
        payload.update(overrides)
        return payload

    def test_pagination_preserves_search_in_both_lists(self):
        for i in range(26):
            Donor.objects.create(name=f'Maria & José {i:02d}')
            Donation.objects.create(donor=self.donor, created_by=self.operator, donation_date=timezone.localdate(), amount='1.00', payment_method='pix')
        for name in ['donors:list', 'donations:list']:
            with self.subTest(name=name):
                response = self.client.get(reverse(name), {'q': 'Maria & José'})
                self.assertContains(response, 'Próxima')
                self.assertContains(response, 'class="pagination-actions"')
                self.assertContains(response, 'page=2')
                self.assertContains(response, 'q=Maria+%26+Jos%C3%A9')
                response = self.client.get(reverse(name), {'q': 'Maria & José', 'page': 2})
                self.assertEqual(response.context['page_obj'].number, 2)
                self.assertEqual(len(response.context['object_list']), 12)
                self.assertContains(response, 'Anterior')

    def test_xlsx_has_typed_cells_and_no_formula_execution(self):
        self.donor.name = '=2+3'
        self.donor.save()
        response = self.client.get(reverse('donations:export-xlsx'))
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(sheet['C2'].value, '=2+3')
        self.assertEqual(sheet['C2'].data_type, 's')
        self.assertEqual(sheet['D2'].value, '52998224725')
        self.assertEqual(sheet['D2'].data_type, 's')
        self.assertEqual(sheet['E2'].value, 1234.56)
        self.assertEqual(sheet['E2'].data_type, 'n')
        self.assertEqual(sheet['F2'].value, 0)
        self.assertEqual(sheet['G2'].value, 1234.56)
        self.assertEqual(sheet['A2'].value.date(), self.donation.donation_date)
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:I2')
        workbook.close()

    def test_exports_apply_filters_and_reject_invalid_filters(self):
        self.assertEqual(self.client.get(reverse('donations:export-xlsx'), {'start_date': 'invalid'}).status_code, 400)
        response = self.client.get(reverse('donations:export-xlsx'), {'payment_method': 'cash'})
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.active.max_row, 1)
        workbook.close()

    def test_cancellation_preserves_history_and_removes_all_financial_totals(self):
        response = self.client.get(self.url('cancel'))
        self.donation.refresh_from_db()
        self.assertEqual(self.donation.status, 'active')
        self.assertContains(response, 'Confirmar cancelamento')
        self.assertEqual(self.client.post(self.url('cancel'), {'cancellation_reason': '   '}).status_code, 200)
        response = self.client.post(self.url('cancel'), {'cancellation_reason': 'Lançamento duplicado'})
        self.assertEqual(response.status_code, 302)
        self.donation.refresh_from_db()
        self.assertEqual(self.donation.status, 'cancelled')
        self.assertEqual(self.donation.cancelled_by, self.admin)
        self.assertIsNotNone(self.donation.cancelled_at)
        self.assertEqual(self.client.post(self.url('cancel'), {'cancellation_reason': 'Repetido'}).status_code, 404)
        self.assertEqual(self.client.post(self.url('update'), self.donation_payload()).status_code, 404)
        for name in ['receipt', 'receipt-download']:
            self.assertEqual(
                self.client.get(
                    reverse(f'donations:{name}', args=[self.donation.pk, 'donation'])
                ).status_code,
                410,
            )
        response = self.client.get(reverse('donations:list'))
        self.assertContains(response, self.donation.receipt_number)
        response = self.client.get(self.url('detail'))
        self.assertContains(response, 'Cancelada')
        response = self.client.get(reverse('donations:list'))
        self.assertEqual(response.context['summary']['total'], Decimal('0'))
        response = self.client.get(reverse('dashboard:index'))
        self.assertEqual(response.context['summary']['month_total'], Decimal('0'))
        self.assertTrue(all(item['total'] == 0 for item in response.context['monthly_chart']))
        response = self.client.get(reverse('donors:detail', args=[self.donor.pk]))
        self.assertEqual(response.context['total_donated'], Decimal('0'))
        self.assertContains(response, 'Cancelada')
        entry = AuditEntry.objects.get(action='cancel')
        self.assertEqual(entry.before['status'], 'active')
        self.assertEqual(entry.after['status'], 'cancelled')
        self.assertEqual(entry.actor, self.admin)

    def test_reader_and_operator_cannot_bypass_access_controls(self):
        self.client.force_login(self.reader)
        for name in ['dashboard:index', 'donations:list', 'donors:list', 'donations:export-xlsx']:
            self.assertEqual(self.client.get(reverse(name)).status_code, 200, name)
        for url in [reverse('donations:create'), self.url('update'), self.url('cancel'), reverse('donors:create'), reverse('donors:update', args=[self.donor.pk])]:
            self.assertEqual(self.client.get(url).status_code, 403, url)
            self.assertEqual(self.client.post(url, self.donation_payload()).status_code, 403, url)
        self.assertEqual(self.client.get(reverse('audit:list')).status_code, 403)
        self.client.force_login(self.operator)
        self.assertEqual(self.client.post(reverse('donations:create'), self.donation_payload()).status_code, 302)
        self.assertEqual(self.client.post(self.url('cancel'), {'cancellation_reason': 'Duplicado'}).status_code, 403)

    def test_unassigned_and_anonymous_users_cannot_access_data(self):
        urls = [reverse('dashboard:index'), reverse('donors:list'), reverse('donations:list'), reverse('donations:export-xlsx'), reverse('donations:receipt', args=[self.donation.pk, 'donation']), reverse('audit:list')]
        self.client.force_login(self.unassigned)
        for url in urls:
            self.assertEqual(self.client.get(url).status_code, 403, url)
        self.client.logout()
        for url in urls:
            self.assertEqual(self.client.get(url).status_code, 302, url)

    def test_cancellation_requires_csrf(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.admin)
        self.assertEqual(client.post(self.url('cancel'), {'cancellation_reason': 'Duplicado'}).status_code, 403)

    def test_audit_failure_rolls_back_the_donation_edit(self):
        with patch('apps.audit.mixins.AuditEntry.objects.create', side_effect=RuntimeError('audit unavailable')):
            with self.assertRaises(RuntimeError):
                self.client.post(self.url('update'), self.donation_payload(donation_amount='999.00'))
        self.donation.refresh_from_db()
        self.assertEqual(self.donation.amount, Decimal('1234.56'))
        self.assertEqual(AuditEntry.objects.count(), 0)

    def test_audit_records_before_after_and_is_read_only(self):
        self.client.post(self.url('update'), self.donation_payload(donation_amount='999.00'))
        entry = AuditEntry.objects.get()
        self.assertEqual(entry.before['valor'], '1234.56')
        self.assertEqual(entry.after['valor'], '999.00')
        self.assertContains(self.client.get(reverse('audit:list'), {'entity': 'donations.donation', 'object_id': self.donation.pk}), '999.00')
        response = self.client.get(self.url('detail'))
        self.assertContains(response, 'Histórico de alterações')
        self.assertContains(response, '999.00')
        self.assertNotContains(response, f'?entity=donations.donation&amp;object_id={self.donation.pk}')
        with self.assertRaises(TypeError):
            entry.save()
        with self.assertRaises(TypeError):
            AuditEntry.objects.all().update(action='cancel')
        with self.assertRaises(TypeError):
            AuditEntry.objects.all().delete()

    def test_setup_roles_configures_the_three_profiles(self):
        call_command('setup_roles', stdout=StringIO())
        self.assertEqual(Group.objects.count(), 3)
        for name in ['Operador', 'Consulta']:
            self.assertFalse(
                Group.objects.get(name=name).permissions.filter(
                    codename__startswith='delete_'
                ).exists()
            )
        self.assertFalse(self.unassigned.groups.exists())
