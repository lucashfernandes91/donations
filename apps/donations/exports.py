from io import BytesIO
import re

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


HEADERS = ['Data', 'Comprovante', 'Doador', 'CPF', 'Doação', 'Contribuição', 'Total', 'Forma de recebimento', 'Observação']


def donation_values(donation) -> list:
    return [
        donation.donation_date, donation.receipt_number, donation.donor.name,
        donation.donor.cpf, donation.donation_amount, donation.contribution_amount,
        donation.amount, donation.get_payment_method_display(), donation.notes,
    ]


def export_xlsx(queryset) -> HttpResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Doações'
    sheet.append(HEADERS)
    for donation in queryset.iterator(chunk_size=500):
        values = donation_values(donation)
        row = sheet.max_row + 1
        for column, value in enumerate(values, start=1):
            if isinstance(value, str):
                value = re.sub(r'[\x00-\x08\x0b-\x0c\x0e-\x1f]', '', value)
            cell = sheet.cell(row, column, value)
            if isinstance(value, str):
                cell.data_type = 's'
            elif column == 1:
                cell.number_format = 'dd/mm/yyyy'
            elif column in (5, 6, 7):
                cell.number_format = '"R$" #,##0.00'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='465FFF')
    for column, width in zip('ABCDEFGHI', [14, 32, 35, 18, 16, 16, 16, 28, 45]):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="doacoes.xlsx"'
    response['Cache-Control'] = 'no-store'
    return response
